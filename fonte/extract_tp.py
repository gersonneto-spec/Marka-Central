"""Extrai a planilha Medição Trem de Passageiros para JSON (motor do painel HTML).

Uso: python3 extract_tp.py "Medição Trem de Passageiros.xlsx" > tp_data.json

Abas lidas:
  QQP                  -> itens contratuais (mestre, tabela TREM_QQP3)
  Insumos por Medição  -> insumos por medição (classe/tipo, qtd e R$ med 1..14)
  OBRAS TP - MED       -> impacto e classificação por medição (EAP + CPU)
  Comparativo Supressão-> preços terceirizado x GRAMAR
  LINHA DO TEMPO ÔNIBUS-> previsto x real por mês
  CPU                  -> BDI da obra
Regra de item folha na QQP: mesma da planilha (UNIDADE <> "").
"""
import sys, json, re, datetime, warnings
warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.utils import column_index_from_string as ci

src = sys.argv[1] if len(sys.argv) > 1 else "Medição Trem de Passageiros.xlsx"
wb = openpyxl.load_workbook(src, data_only=True)


def num(v):
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def txt(v):
    return "" if v is None else str(v).replace("\xa0", " ").strip()


# ---------------------------------------------------------------- QQP
ws = wb["QQP"]
hdr = {c.column: c.value for c in ws[3] if c.value}
bm_cols = []
for col, name in sorted(hdr.items()):
    if isinstance(name, str) and name.startswith("QuantidadeBM"):
        m = re.match(r"BM0*(\d+)(?:,(\d+))?", name[len("Quantidade"):])
        lbl = f"BM{int(m.group(1)):02d}" + (f".{m.group(2)}" if m.group(2) else "")
        bm_cols.append((col, lbl))

base_map = {"Quantidade0": "ct", "Quantidade2": "oin", "Quantidade1": "tac1",
            "Quantidade3": "tac2", "Quantidade4": "tac3", "Quantidade5": "tac4"}
base_cols = {base_map[name]: col for col, name in hdr.items() if name in base_map}

items, groups = [], []
for r in range(4, ws.max_row + 1):
    row = ws[r]
    g = lambda L: row[ci(L) - 1].value
    unit, item, desc = txt(g("P")), txt(g("I")), txt(g("J"))
    if not unit:
        if item or desc:
            groups.append({"r": r, "n1": txt(g("E")), "item": item, "desc": desc})
        continue
    items.append({
        "r": r,
        "contrato": txt(g("A")), "emp": txt(g("B")),
        "n1": txt(g("E")), "n2": txt(g("F")), "n3": txt(g("G")),
        "sgc": txt(g("H")), "item": item, "desc": desc, "cm": txt(g("N")),
        "un": unit, "pu": num(g("R")),
        "base": {k: num(row[c - 1].value) for k, c in base_cols.items()},
        "bm": [num(row[c - 1].value) for c, _ in bm_cols],
    })

# ------------------------------------------------- Insumos por Medição
insumos = []
if "Insumos por Medição" in wb.sheetnames:
    wi = wb["Insumos por Medição"]
    for r in range(4, wi.max_row + 1):
        row = wi[r]
        g = lambda L: row[ci(L) - 1].value
        if not txt(g("D")):
            continue
        insumos.append({
            "classe": txt(g("A")), "tipo": txt(g("B")), "cod": txt(g("C")),
            "desc": txt(g("D")), "un": txt(g("E")),
            "qOrc": num(g("F")), "pu": num(g("G")), "vOrc": num(g("H")),
            "q": [num(row[ci("I") - 1 + i].value) for i in range(14)],
            "v": [num(row[ci("Y") - 1 + i].value) for i in range(14)],
        })

# --------------------------------------------------- OBRAS TP - MED
impacto, meses = [], {}
if "OBRAS TP - MED" in wb.sheetnames:
    wo = wb["OBRAS TP - MED"]
    first = ci("AA")                      # 1ª Medição: QTD, VALOR, IMPACTO, CLASS., %
    for k in range(14):
        mes = txt(wo.cell(row=10, column=first + k * 5 + 1).value)
        if mes:
            meses[f"BM{k + 1:02d}"] = mes
    for r in range(12, wo.max_row + 1):
        row = wo[r]
        g = lambda L: row[ci(L) - 1].value
        if not txt(g("P")) or txt(g("A")) == "TOTAL GERAL":
            continue
        med = []
        for k in range(14):
            b = first + k * 5 - 1
            med.append({"q": num(row[b].value), "v": num(row[b + 1].value),
                        "imp": txt(row[b + 2].value), "cls": txt(row[b + 3].value)})
        impacto.append({
            "bloco": txt(g("A")), "n1": txt(g("B")), "n2": txt(g("C")), "n3": txt(g("D")),
            "n4": txt(g("E")), "item": txt(g("F")), "cpu": txt(g("O")), "desc": txt(g("P")),
            "un": txt(g("Q")), "crit": txt(g("R")), "qtd": num(g("W")), "pu": num(g("X")),
            "total": num(g("Y")), "med": med,
        })

# ------------------------------------------------ Comparativo Supressão
supressao = None
if "Comparativo Supressão" in wb.sheetnames:
    wc = wb["Comparativo Supressão"]
    bdi = num(wb["CPU"]["D7"].value) if "CPU" in wb.sheetnames else 0.0
    linhas = []
    for r, bloco_ini in ((7, 12), (8, 31)):
        linhas.append({
            "desc": txt(wc.cell(row=r, column=2).value), "un": txt(wc.cell(row=r, column=3).value),
            "puComBDI": num(wc.cell(row=r, column=4).value),
            "puSemBDI": num(wc.cell(row=r, column=5).value),
            "puGramar": num(wc.cell(row=r, column=6).value),
            "qtd": [num(wc.cell(row=bloco_ini + k, column=3).value) for k in range(14)],
            "acum": num(wc.cell(row=bloco_ini + 14, column=3).value),
            "saldo": num(wc.cell(row=bloco_ini + 15, column=3).value),
        })
    supressao = {
        "bdi": bdi, "linhas": linhas,
        "mobGramar": num(wc["E68"].value),
        "faturado": [txt(wc.cell(row=51 + k, column=7).value) for k in range(14)],
    }

# ------------------------------------------------ LINHA DO TEMPO ÔNIBUS
onibus = None
for nome in wb.sheetnames:
    if nome.upper().startswith("LINHA DO TEMPO"):
        wl = wb[nome]
        onibus = {
            "precoPrev": num(wl["C6"].value), "custoReal": num(wl["C7"].value),
            "viagemExtra": num(wl["C8"].value),
            "meses": [{"mes": txt(wl.cell(row=12 + k, column=2).value),
                       "qtd": num(wl.cell(row=12 + k, column=3).value),
                       "prev": num(wl.cell(row=12 + k, column=4).value),
                       "real": wl.cell(row=12 + k, column=5).value}
                      for k in range(14)],
        }
        for m in onibus["meses"]:
            m["real"] = None if m["real"] in (None, "") else float(m["real"])
        break

out = {
    "meta": {
        "arquivo": src.split("/")[-1],
        "extraido_em": datetime.date.today().isoformat(),
        "bms": [l for _, l in bm_cols],
        "meses": meses,
        "ref": {"ct": num(ws["S2"].value), "medido": num(ws["U2"].value), "saldo": num(ws["W2"].value)},
        "bdi": supressao["bdi"] if supressao else 0.0,
    },
    "items": items, "groups": groups, "insumos": insumos, "impacto": impacto,
    "supressao": supressao, "onibus": onibus,
}
json.dump(out, sys.stdout, ensure_ascii=False, separators=(",", ":"))
