# -*- coding: utf-8 -*-
"""Extrai as planilhas de DRE e custos da Obra 754 (OCG) para JSON.

Uso: python3 extract_dre.py "DRE_754.xlsx" "Painel_Custos.xlsx" > dre_data.json

Abas lidas:
  DRE_754.xlsx        -> Consolidado (conta x mês, base mestre do DRE)
  Painel_Custos.xlsx  -> EVOLUCAO MENSAL (grupos por mês, inclui ago/26 previsto)
                         MATERIAL COMPARATIVO (fornecedor e categoria, mai a ago/26)
                         Data Sources (procedência das extrações do SIENGE)
"""
import sys, json, datetime, warnings, re
warnings.filterwarnings("ignore")
import openpyxl

src_dre = sys.argv[1] if len(sys.argv) > 1 else "DRE_754.xlsx"
src_cus = sys.argv[2] if len(sys.argv) > 2 else "Painel_Custos.xlsx"


def num(v):
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def txt(v):
    return "" if v is None else str(v).replace("\xa0", " ").strip()


def isnum(v):
    return isinstance(v, (int, float))


# ------------------------------------------------------------ DRE
wb = openpyxl.load_workbook(src_dre, data_only=True)
ws = wb["Consolidado"]
meses = [txt(ws.cell(row=5, column=c).value) for c in range(4, 21)]      # D..T
NM = len(meses)

receita, iss, contas, grupo_atual = None, None, [], None
for r in range(6, 80):
    cod, conta = txt(ws.cell(row=r, column=2).value), txt(ws.cell(row=r, column=3).value)
    if not cod and not conta:
        continue
    vals = [num(ws.cell(row=r, column=c).value) for c in range(4, 4 + NM)]
    if conta.startswith("(=) RECEITA BRUTA"):
        receita = vals; continue
    if "ISS" in conta and cod == "02.01":
        iss = vals; continue
    if conta.startswith("(=) RECEITA LÍQUIDA") or conta.startswith("(=) SOMATÓRIO"):
        continue
    m = re.match(r"^(04\.0\d)\s*·\s*(.+)$", conta)
    if m:
        grupo_atual = {"cod": m.group(1), "nome": m.group(2).strip()}
        continue
    if conta.startswith("SUBTOTAL"):
        continue
    if cod and grupo_atual and cod.startswith(grupo_atual["cod"]):
        contas.append({"cod": cod, "conta": conta, "grupo": grupo_atual["nome"],
                       "gcod": grupo_atual["cod"], "v": [abs(x) for x in vals]})
    elif cod.startswith("06."):
        contas.append({"cod": cod, "conta": conta, "grupo": "CUSTOS FIXOS ADM",
                       "gcod": "06.01", "v": [abs(x) for x in vals]})

# ------------------------------------------------- Painel de custos
wc = openpyxl.load_workbook(src_cus, data_only=True)
we = wc["EVOLUCAO MENSAL"]
serie = []
for r in range(5, 23):
    mes = txt(we.cell(row=r, column=2).value)
    if not mes or mes in ("TOTAL", "MEDIA"):
        continue
    serie.append({"mes": mes,
                  "obra": num(we.cell(row=r, column=3).value), "pessoal": num(we.cell(row=r, column=4).value),
                  "veic": num(we.cell(row=r, column=5).value), "gerais": num(we.cell(row=r, column=6).value),
                  "adm": num(we.cell(row=r, column=7).value), "total": num(we.cell(row=r, column=8).value),
                  "material": num(we.cell(row=r, column=9).value)})
nota_serie = txt(we["B26"].value)

wm = wc["MATERIAL COMPARATIVO"]
mat_meses = [txt(wm.cell(row=9, column=c).value) for c in (4, 6, 12, 17)]   # Maio, Junho, Julho, Agosto
mat_meses = [m.replace("Valor ", "").replace(" (R$)", "") for m in mat_meses]
fornecedores = []
for r in range(10, 86):
    nome = txt(wm.cell(row=r, column=2).value)
    if not nome or nome.startswith("TOTAL"):
        continue
    fornecedores.append({"nome": nome, "cat": txt(wm.cell(row=r, column=3).value),
                         "v": [num(wm.cell(row=r, column=c).value) for c in (4, 6, 12, 17)],
                         "lanc": [int(num(wm.cell(row=r, column=c).value)) for c in (10, 11, 16, 21)]})
categorias = []
for r in range(100, 130):
    nome = txt(wm.cell(row=r, column=2).value)
    if not nome or nome == "TOTAL":
        continue
    categorias.append({"nome": nome, "v": [num(wm.cell(row=r, column=c).value) for c in (3, 5, 9, 13)]})
mat_serie = []
for r in range(91, 97):
    mes = txt(wm.cell(row=r, column=2).value)
    if not mes or mes.startswith("Acum"):
        continue
    mat_serie.append({"mes": mes, "total": num(wm.cell(row=r, column=3).value),
                      "lanc": int(num(wm.cell(row=r, column=6).value))})

fontes = []
wd = wc["Data Sources"]
for r in range(2, wd.max_row + 1):
    if txt(wd.cell(row=r, column=2).value):
        fontes.append({"origem": txt(wd.cell(row=r, column=2).value), "objeto": txt(wd.cell(row=r, column=3).value),
                       "params": txt(wd.cell(row=r, column=6).value), "quando": txt(wd.cell(row=r, column=7).value),
                       "linhas": txt(wd.cell(row=r, column=9).value), "nota": txt(wd.cell(row=r, column=11).value)})

out = {
    "meta": {
        "arquivos": [src_dre.split("/")[-1], src_cus.split("/")[-1]],
        "extraido_em": datetime.date.today().isoformat(),
        "obra": "754", "contrato": "RT-2180KF-G-17251",
        "meses": meses, "nota_serie": nota_serie, "mat_meses": mat_meses,
        "medicao_qqp": {"contrato": 88714693.0852044, "medido": 44494303.86104617, "bm": "BM16"},
    },
    "receita": receita, "iss": iss, "contas": contas,
    "serie": serie, "fornecedores": fornecedores, "categorias": categorias,
    "mat_serie": mat_serie, "fontes": fontes,
}
json.dump(out, sys.stdout, ensure_ascii=False, separators=(",", ":"))
