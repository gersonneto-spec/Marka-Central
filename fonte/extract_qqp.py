"""Extrai a tabela Cadastro_QQP5 (aba QQP) da planilha Medição OCG para JSON.
Uso: python3 extract_qqp.py "Medição OCG.xlsx" > qqp_data.json
Regra de item folha: mesma da planilha (UNIDADE <> "").
"""
import sys, json, re, datetime, warnings
warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.utils import column_index_from_string as ci

src = sys.argv[1]
wb = openpyxl.load_workbook(src, data_only=True)
ws = wb["QQP"]

hdr3 = {c.column: c.value for c in ws[3] if c.value}
# colunas de quantidade dos BMs: cabeçalho "Quantidade<BM>" a partir de AH; nome do BM na linha 2
bm_cols = []
for col, name in sorted(hdr3.items()):
    if isinstance(name, str) and name.startswith("QuantidadeBM"):
        raw = name[len("Quantidade"):]            # ex.: BM04,1  BM010
        m = re.match(r"BM0*(\d+)(?:,(\d+))?", raw)
        label = f"BM{int(m.group(1)):02d}" + (f".{m.group(2)}" if m.group(2) else "")
        bm_cols.append((col, label))

# colunas base contratual (Quantidade0 = CT inicial, 2 = OIN, 1 = TAC1, 3 = TAC2, 4 = TAC3, 5 = TAC4)
base_map = {"Quantidade0": "ct", "Quantidade2": "oin", "Quantidade1": "tac1",
            "Quantidade3": "tac2", "Quantidade4": "tac3", "Quantidade5": "tac4"}
base_cols = {v: col for col, name in hdr3.items() if name in base_map for v in [base_map[name]]}

def num(v):
    if v is None or v == "": return 0.0
    try: return float(v)
    except: return 0.0

def txt(v):
    return "" if v is None else str(v).strip()

items = []
groups = []   # linhas de agrupamento (sem unidade) para reconstruir a hierarquia da QQP
for r in range(4, ws.max_row + 1):
    row = ws[r]
    g = lambda L: row[ci(L) - 1].value
    unit = txt(g("N"))
    item = txt(g("J"))
    desc = txt(g("K"))
    if not unit:
        if item or desc:
            groups.append({"r": r, "ativo": txt(g("A")), "item": item, "desc": desc})
        continue
    bms = [num(row[col - 1].value) for col, _ in bm_cols]
    d = {
        "r": r,
        "ativo": txt(g("A")), "contrato": txt(g("B")), "deflator": txt(g("C")),
        "n1": txt(g("F")), "n2": txt(g("G")), "emp": txt(g("H")),
        "sgc": txt(g("I")), "item": item, "desc": desc, "cm": txt(g("L")),
        "un": unit, "pu": num(g("P")),
        "base": {k: num(row[c - 1].value) for k, c in base_cols.items()},
        "bm": bms,
    }
    items.append(d)

out = {
    "meta": {
        "arquivo": src.split("/")[-1],
        "extraido_em": datetime.date.today().isoformat(),
        "bms": [lbl for _, lbl in bm_cols],
        "ref_planilha": {"ct_geral": num(ws["Q2"].value), "medido": num(ws["S2"].value), "saldo": num(ws["U2"].value)},
    },
    "items": items,
    "groups": groups,
}
json.dump(out, sys.stdout, ensure_ascii=False, separators=(",", ":"))
