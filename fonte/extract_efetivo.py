# -*- coding: utf-8 -*-
"""Extrai as duas planilhas SGC de efetivo (Obras Civis e Trem de Passageiros) para JSON.

Uso: python3 extract_efetivo.py OBRAS_CIVIS.xlsx TREM_PASSAGEIROS.xlsx > efetivo_data.json

O painel mostra nome e matrícula: foi decisão do Coordenador de Obras em 08/10/2026, para
poder clicar num número e ver quem são as pessoas. CPF continua fora, e a página leva
noindex para não ser indexada por buscador. Salário saiu do painel: as planilhas SGC não
trazem essa informação e a estimativa por função só atrapalhava a leitura.
"""
import sys, json, re, unicodedata, datetime as dt, warnings
warnings.filterwarnings("ignore")
import openpyxl

HOJE = dt.date(2026, 10, 7)          # data de atualização das planilhas
ARQ = {"OC": sys.argv[1], "TP": sys.argv[2]}

def txt(v): return "" if v is None else str(v).replace("\xa0", " ").strip()
def dat(v):
    if isinstance(v, dt.datetime): return v.date()
    if isinstance(v, dt.date): return v
    return None
def ym(d): return d.strftime("%Y-%m") if d else None
def iso(d): return d.isoformat() if d else None

def nf(s):
    s = unicodedata.normalize("NFD", s).encode("ascii", "ignore").decode().upper()
    return re.sub(r"\s+", " ", re.sub(r"[^A-Z0-9 .]", " ", s)).strip()

ABR = {"TEC.": "TECNICO", "TEC": "TECNICO", "ENC.": "ENCARREGADO", "ENC": "ENCARREGADO",
       "OP.": "OPERADOR", "OP": "OPERADOR", "AUX.": "AUXILIAR", "AUX": "AUXILIAR",
       "SEG.": "SEGURANCA", "EM": "DE", "EQUIP.": "EQUIPAMENTOS", "MUNK": "MUNCK",
       "ADMINSTRATIVO": "ADMINISTRATIVO", "SUPERVISORA": "SUPERVISOR", "CAMIHAO": "CAMINHAO"}
ROM = r"\s+(I{1,3}|IV|V|VI{0,3}|IX|X)$"
SIN = [("MOTORISTA CAMINHAO MUNCK", "OPERADOR DE CAMINHAO MUNCK"),
       ("OPERADOR CAMINHAO MUNCK", "OPERADOR DE CAMINHAO MUNCK"),
       ("SINALEIRO RIGGER", "SINALEIRO DE RIGGER"),
       ("APRENDIZ DE TECNICO DE SEGURANCA DO TRABALHO", "APRENDIZ TECNICO SEG DO TRABALHO"),
       ("APRENDIZ TECNICO DE SEGURANCA DO TRABALHO", "APRENDIZ TECNICO SEG DO TRABALHO"),
       ("MOTORISTA DE CAMINHAO", "MOTORISTA CAMINHAO"),
       ("TECNICO MEDICAO", "TECNICO DE MEDICAO"),
       ("SUPERVISOR DE MEDICAO", "SUPERVISOR A DE MEDICAO")]

def funcao(cargo):
    """Nome de função comparável entre revisões: a planilha de 07/10 largou o nível romano
    (ENCANADOR I virou ENCANADOR) e mistura abreviações. Sem normalizar, a mesma função
    aparece três vezes no gráfico."""
    s = nf(cargo).replace("(A)", " ").replace("(", " ").replace(")", " ")
    s = " ".join(ABR.get(p, p) for p in s.split())
    while re.search(ROM, s): s = re.sub(ROM, "", s).strip()
    s = re.sub(r"\s+", " ", s).strip()
    for a, b in SIN: s = s.replace(a, b)
    return s

def titulo(s):
    """Nome em caixa de título: a planilha vem toda em maiúscula, que grita na tela."""
    peq = {"de", "da", "do", "das", "dos", "e"}
    return " ".join(p.lower() if p.lower() in peq else p.capitalize() for p in s.split())

def sexo(v):
    s = nf(v)
    if s.startswith("F"): return "F"
    if s.startswith("M"): return "M"          # pega "Mascuino", erro de digitação na origem
    return None

def faixaEt(i):
    if i is None: return None
    for lim, rot in ((25, "18 a 24"), (35, "25 a 34"), (45, "35 a 44"), (55, "45 a 54")):
        if i < lim: return rot
    return "55 ou mais"

def faixaTc(t):
    if t is None: return None
    if t < 0.5: return "até 6 meses"
    if t < 1: return "6 meses a 1 ano"
    if t < 2: return "1 a 2 anos"
    if t < 5: return "2 a 5 anos"
    return "5 anos ou mais"

def cabecalho(w, lin=4):
    return {txt(w.cell(row=lin, column=c).value): c
            for c in range(1, w.max_column + 1) if txt(w.cell(row=lin, column=c).value)}

ativos, saidas, afast = [], [], []
vistos = set()

for obra, arq in ARQ.items():
    wb = openpyxl.load_workbook(arq, data_only=True)

    for aba, apr in (("Colaboradores", False), ("Aprendizes", True)):
        if aba not in wb.sheetnames: continue
        w = wb[aba]; H = cabecalho(w)
        for r in range(5, w.max_row + 1):
            g = lambda k: w.cell(row=r, column=H[k]).value if k in H else None
            mat = txt(g("Matrícula"))
            if not mat: continue
            if (mat, obra) in vistos: continue   # a planilha repete a linha em alguns casos
            vistos.add((mat, obra))
            adm, mob = dat(g("Data de Admissão")), dat(g("Data de mobilização"))
            blo, nas = dat(g("Data de bloqueio do crachá/acesso (Específica)")), dat(g("Data de nascimento"))
            cargo = txt(g("Cargo")) or txt(g("Função")) or "não informado"
            idade = (HOJE - nas).days // 365 if nas else None
            tc = round((HOJE - adm).days / 365.25, 2) if adm else None
            loc = txt(g("Local Atual"))
            ativos.append({
                "mat": mat, "n": titulo(txt(g("Nome"))),
                "o": obra, "c": cargo, "fn": funcao(cargo),
                "m": txt(g("MOD/MOI")) or txt(g("MOI/MOD")) or "não informado",
                "loc": loc,
                # "Não está na obra" contém "na obra": testar pelo negativo, nunca pelo positivo
                "naObra": not re.search(r"n[ãa]o est[áa]|processo de mobiliza", loc, re.I),
                "st": txt(g("Status do colaborador")),
                "adm": iso(adm), "ma": ym(adm), "mob": iso(mob), "mm": ym(mob),
                "dmob": (mob - adm).days if adm and mob else None,
                "tc": tc, "ftc": faixaTc(tc),
                "id": idade, "fx": faixaEt(idade), "sx": sexo(txt(g("Sexo"))),
                "blo": (blo - HOJE).days if blo else None, "dblo": iso(blo),
                "apr": apr, "frente": txt(g("Frente de Serviço")) if apr else "",
            })

    if "Demissionais_Desmob" in wb.sheetnames:
        w = wb["Demissionais_Desmob"]; H = cabecalho(w)
        for r in range(5, w.max_row + 1):
            g = lambda k: w.cell(row=r, column=H[k]).value if k in H else None
            mat = txt(g("Matrícula"))
            if not mat or ("D", mat) in vistos: continue   # o mesmo desligamento consta nas duas planilhas
            vistos.add(("D", mat))
            adm, dem, nas = dat(g("Data de Admissão")), dat(g("Data Demissão")), dat(g("Data de nascimento"))
            cargo = txt(g("Função")) or txt(g("Cargo")) or "não informado"
            # algumas linhas vêm com as colunas deslocadas e trazem CPF no campo Motivo
            mot = txt(g("Motivo"))
            if not re.search(r"[A-Za-zÀ-ÿ]", mot): mot = "não informado"
            tc = round((dem - adm).days / 365.25, 2) if adm and dem else None
            if tc is not None and tc < 0: tc = None        # data de saída anterior à admissão
            saidas.append({"mat": mat, "n": titulo(txt(g("Nome"))), "o": obra,
                           "c": cargo, "fn": funcao(cargo), "mot": mot,
                           "mes": ym(dem), "dt": iso(dem), "ma": ym(adm),
                           "tc": tc, "ftc": faixaTc(tc),
                           "fx": faixaEt((dem - nas).days // 365 if nas and dem else None),
                           "sx": sexo(txt(g("Sexo")))})

    if "Afastados" in wb.sheetnames:
        w = wb["Afastados"]; H = cabecalho(w)
        for r in range(5, w.max_row + 1):
            g = lambda k: w.cell(row=r, column=H[k]).value if k in H else None
            mat = txt(g("Matrícula"))
            if not mat or ("A", mat) in vistos: continue
            vistos.add(("A", mat))
            cargo = txt(g("Função")) or txt(g("Cargo")) or "não informado"
            afast.append({"mat": mat, "n": titulo(txt(g("Nome"))), "o": obra,
                          "c": cargo, "fn": funcao(cargo),
                          "st": txt(g("Status do colaborador")),
                          "m": txt(g("MOI/MOD")) or txt(g("MOD/MOI")) or "não informado"})

json.dump({"meta": {"data": HOJE.isoformat(),
                    "ativos": len(ativos), "saidas": len(saidas), "afastados": len(afast)},
           "ativos": ativos, "saidas": saidas, "afastados": afast},
          sys.stdout, ensure_ascii=False, separators=(",", ":"))
